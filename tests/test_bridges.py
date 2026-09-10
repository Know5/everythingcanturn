"""跨域桥测试：markup→document、raster→document。

这两个桥是「文档域作为目标」的入口：MD/HTML/DOCX/EPUB → PDF，
PNG/多图 → PDF。语义断言（不测排版美观，测信息不丢）：
- 全部块类型在产物 PDF 里可取回
- 多页文档分页正确
- 图片像素保留
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest

pymupdf = pytest.importorskip("pymupdf", reason="需要 pymupdf")

from convcore.pipeline import Converter, convert_file  # noqa: E402

PDF = "application/pdf"


@pytest.fixture(scope="module")
def conv() -> Converter:
    return Converter()


def _pdf_text(path: Path) -> str:
    doc = pymupdf.open(path)
    try:
        return "".join(p.get_text() for p in doc)
    finally:
        doc.close()


# ── markup → document ──


def test_markdown_to_pdf_all_block_types(conv, tmp_path):
    """每种块类型都要在 PDF 里取回：标题/列表/代码/表格/引用。"""
    src = tmp_path / "doc.md"
    src.write_text(
        "# 一级标题\n\n"
        "正文段落，含 English 混排。\n\n"
        "## 二级标题\n\n"
        "- 列表甲\n"
        "- 列表乙\n\n"
        "```python\nprint('hello')\n```\n\n"
        "| 列A | 列B |\n|---|---|\n| 数据1 | 数据2 |\n\n"
        "> 引用一句话\n",
        encoding="utf-8",
    )
    res = convert_file(src, "pdf", out_dir=tmp_path, overwrite=True, converter=conv)
    text = _pdf_text(res.outputs[0].path)
    for probe in ["一级标题", "二级标题", "列表甲", "hello", "数据1", "数据2", "引用一句话"]:
        assert probe in text, f"{probe} 未出现在 PDF 里"


def test_markdown_to_pdf_paginates_long_document(conv, tmp_path):
    """超长文档必须分页——排版器的分页状态机不死循环也不挤在一页。"""
    src = tmp_path / "long.md"
    src.write_text("# 标题\n\n" + "很长的一段内容，需要跨页。セグメント。" * 400, encoding="utf-8")
    res = convert_file(src, "pdf", out_dir=tmp_path, overwrite=True, converter=conv)
    doc = pymupdf.open(res.outputs[0].path)
    try:
        assert doc.page_count >= 2, "长文档挤在一页，分页失败"
    finally:
        doc.close()


def test_html_without_charset_decodes_utf8(conv, tmp_path):
    """无 meta charset 的中文 HTML：lxml 会按 ISO-8859-1 猜，必须先验 UTF-8。

    回归：lxml 的猜测行为把中文全部变成 mojibake，转出的 PDF 是乱码。
    """
    src = tmp_path / "no_decl.html"
    src.write_bytes("<h1>中文标题</h1><p>中文段落。</p>".encode("utf-8"))
    res = convert_file(src, "pdf", out_dir=tmp_path, overwrite=True, converter=conv)
    text = _pdf_text(res.outputs[0].path)
    assert "中文标题" in text and "中文段落" in text


def test_html_single_block_element_to_pdf(conv, tmp_path):
    """单个顶层块元素的 HTML 片段也要能转。

    回归：lxml 片段解析里 <h1>x</h1> 的根就是 h1 本身（文本在 .text 而非子节点），
    _html_blocks 只遍历子节点，把这个根漏了，得到空文档被桥接层以「空文档」拒绝。
    """
    src = tmp_path / "solo.html"
    src.write_bytes("<h1>标题</h1>".encode("utf-8"))
    res = convert_file(src, "pdf", out_dir=tmp_path, overwrite=True, converter=conv)
    text = _pdf_text(res.outputs[0].path)
    assert "标题" in text


def test_epub_chapters_roundtrip_through_pdf(conv, tmp_path):
    """EPUB 的 XHTML 带声明（str 会被 lxml 拒绝），桥路径必须还能走通。"""
    import zipfile

    src = tmp_path / "book.epub"
    xhtml = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        "<html xmlns=\"http://www.w3.org/1999/xhtml\"><head><title>书</title></head>"
        "<body><h1>第一章</h1><p>EPUB 章节内容。</p></body></html>"
    )
    with zipfile.ZipFile(src, "w") as zf:
        # EPUB 规范要求 mimetype 是首个条目——嗅探靠它区分 EPUB 和 ZIP
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip")
        zf.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container version="1.0" '
            'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
            '<rootfile full-path="OEBPS/content.opf" '
            'media-type="application/oebps-package+xml"/></rootfiles></container>',
        )
        zf.writestr(
            "OEBPS/content.opf",
            '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" '
            'version="2.0" unique-identifier="id"><metadata/><manifest><item id="c1" '
            'href="c1.xhtml" media-type="application/xhtml+xml"/></manifest><spine>'
            '<itemref idref="c1"/></spine></package>',
        )
        zf.writestr("OEBPS/c1.xhtml", xhtml)

    res = convert_file(src, "pdf", out_dir=tmp_path, overwrite=True, converter=conv)
    text = _pdf_text(res.outputs[0].path)
    assert "EPUB 章节内容" in text


def test_docx_to_pdf_via_bridge(conv, tmp_path):
    """DOCX → [markup] ⇒ [document] → PDF：桥让 Word 也能出 PDF。"""
    pytest.importorskip("docx", reason="需要 python-docx")
    import docx as docxlib

    src = tmp_path / "t.docx"
    d = docxlib.Document()
    d.add_heading("Word 标题", 1)
    d.add_paragraph("Word 正文段落。")
    d.save(src)

    res = convert_file(src, "pdf", out_dir=tmp_path, overwrite=True, converter=conv)
    text = _pdf_text(res.outputs[0].path)
    assert "Word 标题" in text and "Word 正文段落" in text


# ── raster → document ──


def test_png_to_pdf_embeds_pixels(conv, tmp_path):
    from PIL import Image

    src = tmp_path / "img.png"
    Image.new("RGB", (200, 150), (30, 140, 220)).save(src)
    res = convert_file(src, "pdf", out_dir=tmp_path, overwrite=True, converter=conv)

    doc = pymupdf.open(res.outputs[0].path)
    try:
        assert doc.page_count == 1
        images = doc[0].get_images(full=False)
        assert images, "图片未嵌入 PDF"
        # 尺寸映射：200px@96dpi → 150pt
        assert abs(doc[0].rect.width - 150.0) < 1.0
    finally:
        doc.close()


def test_multi_frame_gif_to_pdf_one_page_per_frame(conv, tmp_path):
    """动图 → PDF：每帧一页（时间轴不保留，记 policy 告警）。"""
    from PIL import Image

    src = tmp_path / "anim.gif"
    a = Image.new("RGB", (64, 64), "red")
    b = Image.new("RGB", (64, 64), "blue")
    a.save(src, save_all=True, append_images=[b], duration=100)

    res = convert_file(src, "pdf", out_dir=tmp_path, overwrite=True, converter=conv)
    doc = pymupdf.open(res.outputs[0].path)
    try:
        assert doc.page_count == 2, "每帧一页的排列约定被破坏"
    finally:
        doc.close()
    assert any("帧" in str(w) for w in res.warnings), "时间轴丢弃必须显式告警"


# ── 图拓扑 ──


def test_bridges_visible_in_path_search(conv):
    """新桥必须在 path/targets 里可见——图声明和实现同步的回归。"""
    from convcore.formats import domain_of

    assert domain_of("text/markdown") is not None
    targets = conv.targets_for("text/markdown")
    assert PDF in targets, "MD→PDF 不可达：BRIDGES 表或 _bridge 分发缺失"
    targets = conv.targets_for("image/png")
    assert PDF in targets, "PNG→PDF 不可达"
    # DOCX 经 markup hub 也可达
    DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert PDF in conv.targets_for(DOCX)


def test_svg_to_pdf_through_double_bridge(conv, tmp_path):
    """SVG → [vector] ⇒ [raster] ⇒ [document] → PDF：两跳桥的复合。"""
    pytest.importorskip("resvg_py", reason="需要 resvg-py")
    src = tmp_path / "icon.svg"
    src.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="80" height="80">'
        '<rect width="80" height="80" fill="green"/></svg>',
        encoding="utf-8",
    )
    res = convert_file(src, "pdf", out_dir=tmp_path, overwrite=True, converter=conv)
    doc = pymupdf.open(res.outputs[0].path)
    try:
        assert doc[0].get_images(full=False), "SVG 栅格化后未嵌入 PDF"
    finally:
        doc.close()
