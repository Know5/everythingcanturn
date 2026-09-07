"""数据格式测试：树形（JSON/YAML/TOML/XML）与表格（CSV/TSV/JSONL/XLSX）。

重点是两类真实事故：类型推断毁数据、编码嗅探失败。
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest

from convcore import ConversionError, LossTier, convert_file

pytest.importorskip("yaml")
pytest.importorskip("openpyxl")

RECORDS = [
    {
        "编号": "007",
        "姓名": "张三",
        "身份证": "110101199003078888",
        "工资": 12500.5,
        "在职": True,
    },
    {
        "编号": "042",
        "姓名": "李四",
        "身份证": "310104198812125678",
        "工资": 9800.0,
        "在职": False,
    },
]


@pytest.fixture
def json_file(tmp_path) -> Path:
    p = tmp_path / "staff.json"
    p.write_text(json.dumps(RECORDS, ensure_ascii=False, indent=2), encoding="utf-8")
    return p


@pytest.fixture
def yaml_file(tmp_path) -> Path:
    p = tmp_path / "cfg.yaml"
    p.write_text(
        "name: convcore\n"
        "version: 0.1.0\n"
        "features:\n  - images\n  - media\n"
        "nested:\n  deep:\n    value: 42\n"
        "empty: null\n",
        encoding="utf-8",
    )
    return p


# ── 类型推断：数据转换最经典的事故来源 ──


def test_leading_zeros_survive_json_to_csv(json_file, tmp_path):
    """把 "007" 判成 int 会丢前导零。邮编、编号、电话都会被毁。"""
    res = convert_file(json_file, "csv", out_dir=tmp_path, overwrite=True)
    text = res.outputs[0].path.read_bytes().decode("utf-8-sig")

    rows = list(csv.DictReader(text.splitlines()))
    assert rows[0]["编号"] == "007", f"前导零丢了：{rows[0]['编号']}"
    assert rows[1]["编号"] == "042"


def test_long_digits_not_turned_into_float(json_file, tmp_path):
    """18 位身份证判成 float 会变科学计数法且丢精度，且不可逆。"""
    res = convert_file(json_file, "csv", out_dir=tmp_path, overwrite=True)
    text = res.outputs[0].path.read_bytes().decode("utf-8-sig")

    assert "110101199003078888" in text
    assert "1.101011990030788e" not in text.lower()


def test_type_inference_on_csv_read(tmp_path):
    """CSV 全是字符串，转 JSON 时数字该变数字，但危险的不变。"""
    src = tmp_path / "t.csv"
    src.write_text(
        "编号,数量,单价,启用\n007,100,12.5,true\n012,250,8.0,false\n",
        encoding="utf-8",
    )
    res = convert_file(src, "json", out_dir=tmp_path, overwrite=True)
    data = json.loads(res.outputs[0].path.read_text(encoding="utf-8"))

    assert data[0]["数量"] == 100 and isinstance(data[0]["数量"], int)
    assert data[0]["单价"] == 12.5 and isinstance(data[0]["单价"], float)
    assert data[0]["启用"] is True
    # 前导零列必须保持字符串
    assert data[0]["编号"] == "007"


def test_infer_types_can_be_disabled(tmp_path):
    src = tmp_path / "t.csv"
    src.write_text("a,b\n1,2\n", encoding="utf-8")
    res = convert_file(
        src, "json", out_dir=tmp_path, overwrite=True,
        overrides={"infer_types": False},
    )
    data = json.loads(res.outputs[0].path.read_text(encoding="utf-8"))
    assert data[0]["a"] == "1", "关掉推断后应保持字符串"


def test_type_protection_is_reported(tmp_path):
    """告警只在真正做了推断的路径上出现。

    CSV 读入时一切都是字符串，需要推断，所以会告警。而 JSON 的值本来就带类型，
    走 tree→table 桥接不做推断，那条路径上不该有这个告警。
    """
    src = tmp_path / "codes.csv"
    src.write_text("编号,数量\n007,100\n042,250\n", encoding="utf-8")

    res = convert_file(src, "json", out_dir=tmp_path, overwrite=True)
    assert any("类型保护" in w or "前导零" in w for w in res.warnings), res.warnings

    # 对照：JSON 源不做推断，不该有此告警
    j = tmp_path / "x.json"
    j.write_text(json.dumps([{"编号": "007"}]), encoding="utf-8")
    res2 = convert_file(j, "csv", out_dir=tmp_path, overwrite=True)
    assert not any("类型保护" in w for w in res2.warnings)


# ── 编码 ──


def test_gbk_csv_detected(tmp_path):
    """中文 Excel 默认导出 GBK，按 UTF-8 硬解会乱码。"""
    src = tmp_path / "gbk.csv"
    with src.open("w", encoding="gb18030", newline="") as f:
        w = csv.writer(f)
        w.writerow(["名称", "数量"])
        w.writerow(["螺丝", 100])

    res = convert_file(src, "json", out_dir=tmp_path, overwrite=True)
    data = json.loads(res.outputs[0].path.read_text(encoding="utf-8"))
    assert data[0]["名称"] == "螺丝"
    assert any("编码" in w for w in res.warnings)


def test_utf8_bom_csv_column_names_clean(tmp_path):
    """Excel 的 UTF-8 CSV 带 BOM。不处理会让第一个列名变成 ﻿姓名，取列全失败。"""
    src = tmp_path / "bom.csv"
    src.write_bytes(b"\xef\xbb\xbf" + "姓名,年龄\n张三,30\n".encode())

    res = convert_file(src, "json", out_dir=tmp_path, overwrite=True)
    data = json.loads(res.outputs[0].path.read_text(encoding="utf-8"))
    assert "姓名" in data[0], f"列名被 BOM 污染：{list(data[0])}"


def test_csv_output_has_bom_for_excel(json_file, tmp_path):
    """Excel 打开无 BOM 的 UTF-8 CSV 会把中文显示成乱码。"""
    res = convert_file(json_file, "csv", out_dir=tmp_path, overwrite=True)
    assert res.outputs[0].path.read_bytes().startswith(b"\xef\xbb\xbf")


def test_bom_can_be_disabled(json_file, tmp_path):
    res = convert_file(
        json_file, "csv", out_dir=tmp_path, overwrite=True,
        overrides={"bom": False},
    )
    assert not res.outputs[0].path.read_bytes().startswith(b"\xef\xbb\xbf")


def test_delimiter_sniffed(tmp_path):
    src = tmp_path / "semi.csv"
    src.write_text("a;b;c\n1;2;3\n", encoding="utf-8")
    res = convert_file(src, "json", out_dir=tmp_path, overwrite=True)
    data = json.loads(res.outputs[0].path.read_text(encoding="utf-8"))
    assert list(data[0]) == ["a", "b", "c"], "分号分隔符没嗅探出来"


# ── 树形格式互转 ──


def test_yaml_to_json_lossless(yaml_file, tmp_path):
    import yaml

    res = convert_file(yaml_file, "json", out_dir=tmp_path, overwrite=True)
    assert res.loss is LossTier.LOSSLESS

    src = yaml.safe_load(yaml_file.read_text(encoding="utf-8"))
    got = json.loads(res.outputs[0].path.read_text(encoding="utf-8"))
    assert got == src, "YAML→JSON 应完全等价"


def test_roundtrip_json_yaml_json(json_file, tmp_path):
    y = convert_file(
        json_file, "yaml", out_dir=tmp_path / "a", overwrite=True
    ).outputs[0].path
    back = convert_file(
        y, "json", out_dir=tmp_path / "b", overwrite=True
    ).outputs[0].path

    assert json.loads(back.read_text(encoding="utf-8")) == RECORDS


def test_toml_drops_null_with_warning(yaml_file, tmp_path):
    """TOML 规范没有 null。必须告知，不能静默写成空字符串。"""
    import tomllib

    res = convert_file(yaml_file, "toml", out_dir=tmp_path, overwrite=True)
    assert res.loss is LossTier.STRUCTURAL
    assert any("null" in w or "空值" in w for w in res.warnings)

    data = tomllib.loads(res.outputs[0].path.read_text(encoding="utf-8"))
    assert "empty" not in data, "null 键应被省略而非写成空值"
    assert data["nested"]["deep"]["value"] == 42, "嵌套结构应保留"


def test_toml_wraps_top_level_array(json_file, tmp_path):
    """TOML 顶层必须是表，数组要包一层且必须告知键名是我们加的。"""
    import tomllib

    res = convert_file(json_file, "toml", out_dir=tmp_path, overwrite=True)
    assert any("顶层" in w for w in res.warnings)

    data = tomllib.loads(res.outputs[0].path.read_text(encoding="utf-8"))
    assert "items" in data
    assert len(data["items"]) == 2


def test_xml_roundtrip_preserves_structure(tmp_path):
    """XML→树→XML 应保住根标签名和层级。"""
    src = tmp_path / "a.xml"
    src.write_text(
        '<config version="2"><name>test</name>'
        "<items><item>1</item><item>2</item></items></config>",
        encoding="utf-8",
    )
    j = convert_file(src, "json", out_dir=tmp_path / "a", overwrite=True).outputs[0].path
    data = json.loads(j.read_text(encoding="utf-8"))

    assert "config" in data, "根标签名丢了"
    assert data["config"]["@version"] == 2, "属性应映射为 @name"
    assert data["config"]["items"]["item"] == [1, 2], "重复标签应合并为数组"

    back = convert_file(j, "xml", out_dir=tmp_path / "b", overwrite=True).outputs[0].path
    xml = back.read_text(encoding="utf-8")
    assert xml.lstrip().startswith("<config"), f"根标签没还原：{xml[:60]}"
    assert 'version="2"' in xml


def test_xml_sanitizes_illegal_tag_names(tmp_path):
    """JSON 键可以是任意字符串，XML 标签不能以数字开头也不能含空格。"""
    src = tmp_path / "bad.json"
    src.write_text(
        json.dumps({"2020 年": 1, "a b c": 2, "正常": 3}, ensure_ascii=False),
        encoding="utf-8",
    )
    res = convert_file(src, "xml", out_dir=tmp_path, overwrite=True)

    from lxml import etree

    # 能解析就说明标签名合法
    root = etree.fromstring(res.outputs[0].path.read_bytes())
    tags = [c.tag for c in root]
    assert all(not t[0].isdigit() for t in tags), f"标签以数字开头：{tags}"
    assert all(" " not in t for t in tags)


def test_json_parse_error_reports_location(tmp_path):
    src = tmp_path / "bad.json"
    src.write_text('{"a": 1,\n "b": }', encoding="utf-8")
    with pytest.raises(ConversionError, match="第 2 行"):
        convert_file(src, "yaml", out_dir=tmp_path, overwrite=True)


# ── 表格 ──


def test_xlsx_roundtrip(json_file, tmp_path):
    x = convert_file(
        json_file, "xlsx", out_dir=tmp_path / "a", overwrite=True
    ).outputs[0].path
    back = convert_file(
        x, "json", out_dir=tmp_path / "b", overwrite=True
    ).outputs[0].path
    data = json.loads(back.read_text(encoding="utf-8"))

    assert len(data) == 2
    assert data[0]["姓名"] == "张三"
    assert data[0]["编号"] == "007", "Excel 往返后前导零丢了"


def test_xlsx_has_bold_header_and_freeze(json_file, tmp_path):
    import openpyxl

    out = convert_file(
        json_file, "xlsx", out_dir=tmp_path, overwrite=True
    ).outputs[0].path
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert ws["A1"].font.bold, "表头应加粗"
    assert ws.freeze_panes == "A2", "表头应冻结"


def test_jsonl_roundtrip(json_file, tmp_path):
    jl = convert_file(
        json_file, "jsonl", out_dir=tmp_path / "a", overwrite=True
    ).outputs[0].path
    lines = [l for l in jl.read_text(encoding="utf-8").splitlines() if l.strip()]
    assert len(lines) == 2
    assert json.loads(lines[0])["姓名"] == "张三"


def test_tsv_uses_tabs(json_file, tmp_path):
    res = convert_file(json_file, "tsv", out_dir=tmp_path, overwrite=True)
    text = res.outputs[0].path.read_bytes().decode("utf-8-sig")
    assert "\t" in text.splitlines()[0]


def test_multi_sheet_xlsx_becomes_multiple_csv(tmp_path):
    """单表格式遇到多表必须一进多出。静默只写第一张是数据丢失。"""
    import openpyxl

    src = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in ("一月", "二月", "三月"):
        ws = wb.create_sheet(name)
        ws.append(["项目", "金额"])
        ws.append(["A", 100])
    wb.save(src)

    res = convert_file(src, "csv", out_dir=tmp_path, overwrite=True)
    assert len(res.outputs) == 3, f"3 张表应输出 3 个文件，实为 {len(res.outputs)}"
    assert any("多工作表" in w for w in res.warnings)
    names = [o.path.name for o in res.outputs]
    assert len(set(names)) == 3, f"文件名重复：{names}"


def test_sheet_selection(tmp_path):
    import openpyxl

    src = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in ("Alpha", "Beta"):
        ws = wb.create_sheet(name)
        ws.append(["k", "v"])
        ws.append([name, 1])
    wb.save(src)

    res = convert_file(
        src, "json", out_dir=tmp_path, overwrite=True, overrides={"sheet": "Beta"}
    )
    data = json.loads(res.outputs[0].path.read_text(encoding="utf-8"))
    assert data[0]["k"] == "Beta"


def test_unknown_sheet_name_lists_options(tmp_path):
    import openpyxl

    src = tmp_path / "m.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Real"
    wb.active.append(["a"])
    wb.save(src)

    with pytest.raises(ConversionError, match="Real"):
        convert_file(
            src, "csv", out_dir=tmp_path, overwrite=True,
            overrides={"sheet": "Nonexistent"},
        )


# ── 跨域桥接 ──


def test_tree_to_table_rejects_wrong_shape(tmp_path):
    """标量数组变不成表格。必须明确说需要什么形状。"""
    src = tmp_path / "arr.json"
    src.write_text("[1, 2, 3]", encoding="utf-8")

    with pytest.raises(ConversionError, match="对象数组"):
        convert_file(src, "csv", out_dir=tmp_path, overwrite=True)


def test_tree_to_table_serializes_nested(tmp_path):
    """嵌套值序列化为 JSON 字符串，而不是压平成 a.b.c 列名。"""
    src = tmp_path / "n.json"
    src.write_text(
        json.dumps([{"id": 1, "tags": ["x", "y"], "meta": {"k": "v"}}]),
        encoding="utf-8",
    )
    res = convert_file(src, "csv", out_dir=tmp_path, overwrite=True)
    text = res.outputs[0].path.read_bytes().decode("utf-8-sig")

    assert any("嵌套" in w for w in res.warnings)
    assert "a.b" not in text, "不该压平成点号列名"
    rows = list(csv.DictReader(text.splitlines()))
    assert json.loads(rows[0]["tags"]) == ["x", "y"]


def test_csv_to_markdown_table(tmp_path):
    """CSV → Markdown 表格：把数据变成人看的东西。"""
    src = tmp_path / "d.csv"
    src.write_text("名称,数量\n螺丝,100\n螺母,250\n", encoding="utf-8")
    res = convert_file(src, "md", out_dir=tmp_path, overwrite=True)
    text = res.outputs[0].path.read_text(encoding="utf-8")

    assert "| 名称 | 数量 |" in text
    assert "| --- |" in text
    assert "| 螺丝 | 100 |" in text


def test_xlsx_to_html_renders_table(json_file, tmp_path):
    x = convert_file(
        json_file, "xlsx", out_dir=tmp_path / "a", overwrite=True
    ).outputs[0].path
    res = convert_file(x, "html", out_dir=tmp_path / "b", overwrite=True)
    html = res.outputs[0].path.read_text(encoding="utf-8")

    assert "<table>" in html and "<th>" in html
    assert "张三" in html
