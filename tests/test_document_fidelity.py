"""文档域保真修复的回归测试（审计清单驱动）。

覆盖：Markdown 相对路径图片进 PDF、缺图占位与告警、长单元格不截断、
colspan 合并、rowspan 告警、多表提取与多 sheet 导出。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image

from convcore.pipeline import convert_file

pytest.importorskip("pymupdf", reason="文档桥需要 pymupdf")


def _pdf_text(path: Path) -> str:
    import pymupdf

    with pymupdf.open(path) as doc:
        return "".join(p.get_text() for p in doc)


def _squeeze(text: str) -> str:
    return "".join(text.split())


def test_markdown_image_is_embedded_in_pdf(tmp_path):
    """相对路径的图片要按源文件目录解析并真正嵌进 PDF，而不是留占位。"""
    Image.new("RGB", (120, 60), "blue").save(tmp_path / "pic.png")
    md = tmp_path / "doc.md"
    md.write_text("# 标题\n\n![蓝色的图](pic.png)\n\n尾段。\n", encoding="utf-8")

    res = convert_file(md, "pdf", out_dir=tmp_path, overwrite=True)
    import pymupdf

    with pymupdf.open(res.outputs[0].path) as doc:
        assert any(page.get_images() for page in doc), "图片没有被嵌入 PDF"
    assert not any("占位" in str(w) for w in res.warnings), "图片应该被找到并嵌入"
    assert "尾段" in _squeeze(_pdf_text(res.outputs[0].path))


def test_missing_image_leaves_searchable_placeholder(tmp_path):
    """图片找不到时留可检索的占位 + DROP 告警，不许静默消失。"""
    md = tmp_path / "doc.md"
    md.write_text("# T\n\n![缺失](nope.png)\n", encoding="utf-8")
    res = convert_file(md, "pdf", out_dir=tmp_path, overwrite=True)
    assert any("图片" in str(w) for w in res.warnings), "缺图必须有告警"
    assert "nope.png" in _squeeze(_pdf_text(res.outputs[0].path))


def test_long_table_cell_is_not_truncated(tmp_path):
    """行高必须随内容换行——固定行高会把长单元格直接裁掉。"""
    long = "很长的中文单元格内容" * 10
    md = tmp_path / "t.md"
    md.write_text(f"| A | B |\n|---|---|\n| {long} | x |\n", encoding="utf-8")
    res = convert_file(md, "pdf", out_dir=tmp_path, overwrite=True)
    text = _squeeze(_pdf_text(res.outputs[0].path))
    assert _squeeze(long)[-15:] in text, "长单元格被固定行高裁剪"


def test_colspan_table_renders_merged_cell(tmp_path):
    html = tmp_path / "t.html"
    html.write_text(
        '<table><tr><td colspan="2">合并标题</td></tr>'
        "<tr><td>a1</td><td>b1</td></tr></table>",
        encoding="utf-8",
    )
    res = convert_file(html, "pdf", out_dir=tmp_path, overwrite=True)
    text = _squeeze(_pdf_text(res.outputs[0].path))
    assert "合并标题" in text and "a1" in text and "b1" in text


def test_rowspan_table_warns_instead_of_misplacing(tmp_path):
    html = tmp_path / "t.html"
    html.write_text(
        '<table><tr><td rowspan="2">跨行</td><td>a</td></tr>'
        "<tr><td>b</td></tr></table>",
        encoding="utf-8",
    )
    res = convert_file(html, "pdf", out_dir=tmp_path, overwrite=True)
    assert any("rowspan" in str(w) for w in res.warnings)


def test_pdf_extraction_keeps_every_table(tmp_path):
    """跨页的多张表全部提取进 Table IR，而不是只留第一张。"""
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import PageBreak, SimpleDocTemplate
    from reportlab.platypus import Table as RLTable

    from convcore.plugins.document import analyze
    from convcore.plugins.document_extract import document_to_table
    from convcore.policy import Policy
    from convcore.types import Resource

    pdf = tmp_path / "two.pdf"
    SimpleDocTemplate(str(pdf), pagesize=A4).build(
        [
            RLTable([["a", "b"], ["1", "2"]]),
            PageBreak(),
            RLTable([["c", "d"], ["3", "4"]]),
        ]
    )
    ref = analyze(Resource(media_type="application/pdf", path=pdf), tmp_path)
    table = document_to_table(ref, {}, Policy())
    assert len(table.sheets) >= 2, f"多表被并成一张：{[s.name for s in table.sheets]}"


def test_multi_sheet_table_pdf_exports_all(tmp_path):
    """table→PDF 多 sheet 全部导出（每张带 sheet 名），不是只排首表。"""
    from convcore.ir.data import Sheet, Table
    from convcore.plugins.document_build import table_to_document
    from convcore.policy import Policy

    table = Table(
        sheets=[
            Sheet(columns=["x"], rows=[["1"]], name="第一表"),
            Sheet(columns=["y"], rows=[["2"]], name="第二表"),
        ]
    )
    ref = table_to_document(table, {}, Policy())
    assert ref.source.path is not None
    text = _squeeze(_pdf_text(ref.source.path))
    assert "第一表" in text and "第二表" in text


def test_pdf_to_markdown_keeps_tables(tmp_path):
    """PDF→Markdown 保留表格结构，而不是把整表压成段落文本。"""
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate
    from reportlab.platypus import Table as RLTable

    # reportlab 默认字体没有 CJK 字形，样本用 ASCII
    pdf = tmp_path / "t.pdf"
    SimpleDocTemplate(str(pdf), pagesize=A4).build(
        [RLTable([["name", "qty"], ["apple", "3"], ["banana", "5"]])]
    )
    res = convert_file(pdf, "md", out_dir=tmp_path, overwrite=True)
    text = res.outputs[0].path.read_text(encoding="utf-8")
    assert "|" in text and "apple" in text, f"表格没有保留：{text[:200]}"


def test_prose_pdf_is_not_mistaken_for_a_table(tmp_path):
    """无线框的散文 PDF 不能被 text 策略误判成表格。"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    styles = getSampleStyleSheet()
    pdf = tmp_path / "prose.pdf"
    SimpleDocTemplate(str(pdf), pagesize=A4).build(
        [
            Paragraph(
                "This is a normal paragraph with several words in it.",
                styles["BodyText"],
            ),
            Paragraph(
                "Another paragraph follows, also plain prose without columns.",
                styles["BodyText"],
            ),
        ]
    )
    res = convert_file(pdf, "md", out_dir=tmp_path, overwrite=True)
    text = res.outputs[0].path.read_text(encoding="utf-8")
    assert "|" not in text, f"散文被误判成表格：{text[:200]}"


def test_pdf_to_markdown_keeps_images(tmp_path):
    """PDF 内嵌图应进入 Markdown，而不是只留一句「图片未保留」。"""
    import pymupdf
    from PIL import Image
    from convcore import convert_file

    image = tmp_path / "red.png"
    Image.new("RGB", (32, 24), "red").save(image)
    pdf = tmp_path / "with-image.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "image page")
    page.insert_image(pymupdf.Rect(72, 100, 180, 180), filename=str(image))
    doc.save(pdf)
    doc.close()

    res = convert_file(pdf, "md", out_dir=tmp_path, overwrite=True)
    text = res.outputs[0].path.read_text(encoding="utf-8")
    assert "![PDF image](data:image/" in text


def test_docx_writer_embeds_data_uri_images(tmp_path):
    """Markdown/HTML 里的 data URI 图片写进 DOCX 时应真正嵌入，而不是占位符。"""
    import base64
    import io
    import zipfile

    from PIL import Image

    from convcore.ir.markup import IMAGE, Block, DocTree
    from convcore.plugins.office import DocxCodec

    buf = io.BytesIO()
    Image.new("RGB", (80, 60), "red").save(buf, "PNG")
    src = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
    docx = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    resource = DocxCodec().write(
        DocTree(blocks=[Block(IMAGE, src=src, alt="red")]), docx, {}
    )[0]
    path = tmp_path / "image.docx"
    path.write_bytes(resource.data or b"")
    with zipfile.ZipFile(path) as archive:
        assert any(n.startswith("word/media/") for n in archive.namelist())


def test_docx_parser_extracts_inline_images(tmp_path):
    """DOCX 里的内嵌图应进入 DocTree，后续才能转 Markdown/HTML。"""
    import docx
    from PIL import Image

    from convcore.ir.markup import IMAGE
    from convcore.plugins.office import DocxCodec
    from convcore.types import Resource

    image = tmp_path / "red.png"
    Image.new("RGB", (24, 18), "red").save(image)
    document = docx.Document()
    document.add_paragraph("before")
    document.add_picture(str(image))
    document.add_paragraph("after")
    path = tmp_path / "images.docx"
    document.save(path)

    tree = DocxCodec().parse(
        [Resource(media_type=DocxCodec.reads[0].media_type, path=path)], {}
    )
    images = [b for b in tree.blocks if b.kind == IMAGE]
    assert len(images) == 1
    assert images[0].src and images[0].src.startswith("data:image/png;base64,")


def test_epub_writer_packages_images(tmp_path):
    """data URI 图片应写进 EPUB 的 OEBPS/images 并进入 manifest。"""
    import base64
    import io
    import zipfile

    from PIL import Image

    from convcore.ir.markup import HEADING, IMAGE, Block, DocTree, Span
    from convcore.plugins.office import EpubCodec

    buf = io.BytesIO()
    Image.new("RGB", (32, 24), "red").save(buf, "PNG")
    src = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
    tree = DocTree(
        blocks=[
            Block(HEADING, spans=[Span("Title")], level=1),
            Block(IMAGE, src=src, alt="red"),
        ]
    )
    resource = EpubCodec().write(tree, "application/epub+zip", {})[0]
    path = tmp_path / "book.epub"
    path.write_bytes(resource.data or b"")
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        assert "OEBPS/images/img001.png" in names
        assert b'media-type="image/png"' in archive.read("OEBPS/content.opf")
        chapter = archive.read("OEBPS/ch001.xhtml").decode("utf-8")
        assert 'src="images/img001.png"' in chapter
