"""文档域转换测试。

重点在 PDF→Word：这是唯一的 INFERRED 等级转换，最容易出错也最需要诚实。
"""

from __future__ import annotations

from pathlib import Path

import pytest

from convcore import ConversionError, LossTier, convert_file

pytest.importorskip("pymupdf", reason="需要 pymupdf")

from .make_pdf import make_plain, make_report, make_scanned  # noqa: E402

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture(scope="module")
def pdfs(tmp_path_factory) -> dict[str, Path]:
    d = tmp_path_factory.mktemp("pdfs")
    return {
        "report": make_report(d / "report.pdf"),
        "scanned": make_scanned(d / "scanned.pdf"),
        "plain": make_plain(d / "plain.pdf"),
    }


# ── 探测 ──


def test_analyze_detects_structure(pdfs):
    from convcore.plugins.document import analyze
    from convcore.types import Resource

    ref = analyze(
        Resource(media_type="application/pdf", path=pdfs["report"]),
        pdfs["report"].parent,
    )
    assert ref.page_count == 2
    assert ref.has_text_layer
    assert not ref.is_scanned_doc
    assert ref.table_count >= 1, "第 1 页的表格没被检测到"
    assert ref.pages[0].char_count > 100


def test_analyze_detects_scanned(pdfs):
    """扫描件识别是转 Word 前最关键的判断。

    有文本层 → 版面推断；无文本层 → 只能得到图片，必须先 OCR。
    这两条路的失败模式完全不同，混在一起处理必然出错。
    """
    from convcore.plugins.document import analyze
    from convcore.types import Resource

    ref = analyze(
        Resource(media_type="application/pdf", path=pdfs["scanned"]),
        pdfs["scanned"].parent,
    )
    assert not ref.has_text_layer
    assert ref.is_scanned_doc
    assert len(ref.scanned_pages) == ref.page_count


def test_page_spec_parsing():
    from convcore.plugins.document import parse_page_spec

    assert parse_page_spec("1-3", 10) == [0, 1, 2]
    assert parse_page_spec("1,3,5", 10) == [0, 2, 4]
    assert parse_page_spec("2", 10) == [1]
    assert parse_page_spec("8-", 10) == [7, 8, 9]
    assert parse_page_spec("-2", 10) == [0, 1]
    assert parse_page_spec(None, 3) == [0, 1, 2]
    assert parse_page_spec([1, 2], 5) == [0, 1]
    # 去重且保序
    assert parse_page_spec("2,1,2", 5) == [1, 0]


def test_page_spec_rejects_out_of_range():
    from convcore.plugins.document import parse_page_spec

    with pytest.raises(ConversionError, match="超出范围"):
        parse_page_spec("5", 3)


# ── PDF → Word：核心功能 ──


def _docx(path: Path):
    import docx

    return docx.Document(str(path))


def test_pdf_to_word_preserves_text_and_table(pdfs, tmp_path):
    """版面还原的基本要求：文字不丢，表格结构还原。"""
    res = convert_file(pdfs["report"], "docx", out_dir=tmp_path, overwrite=True)

    assert res.loss is LossTier.INFERRED, (
        f"PDF→Word 必须是 INFERRED（推断而非翻译），实为 {res.loss.name}"
    )
    assert not res.loss.is_reversible

    d = _docx(res.outputs[0].path)
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Quarterly Report" in text
    assert "Revenue by region" in text

    assert d.tables, "表格完全没还原出来"
    tb = d.tables[0]
    assert len(tb.rows) == 5 and len(tb.columns) == 4, (
        f"表格尺寸错了：{len(tb.rows)}×{len(tb.columns)}，应为 5×4"
    )
    cells = [c.text.strip() for c in tb.rows[0].cells]
    assert cells == ["Region", "Q1", "Q2", "Q3"], f"表头错了：{cells}"
    # 数值不该在还原中被改动
    assert "1,204" in tb.rows[1].cells[1].text


def test_pdf_to_word_warns_it_is_inference(pdfs, tmp_path):
    """必须明确告知这是推断，不是无损转换。

    用户以为能转回去是最常见的误解，静默转换等于误导。
    """
    res = convert_file(pdfs["report"], "docx", out_dir=tmp_path, overwrite=True)
    joined = " ".join(res.warnings)
    assert "推断" in joined
    assert "转回 PDF" in joined


def test_review_marks_inserted_in_document(pdfs, tmp_path):
    """复核提示要写进 Word 文件本身。

    Word 会被转发、归档、打印，脱离转换工具的上下文。把「这是机器推断的结果」
    写进文件，下一个打开它的人才知道不能当原件用。
    """
    res = convert_file(pdfs["report"], "docx", out_dir=tmp_path, overwrite=True)
    d = _docx(res.outputs[0].path)

    head = [p.text.strip() for p in d.paragraphs[:4] if p.text.strip()]
    assert head, "文档开头没有内容"
    assert "自动转换" in head[0], f"第一段应是主提示，实为：{head[0][:40]}"

    # 提示必须有颜色标记，否则和正文混在一起看不出来
    first = d.paragraphs[0]
    assert first.runs
    rgb = first.runs[0].font.color.rgb
    assert rgb is not None and str(rgb) == "996B00"


def test_review_marks_can_be_disabled(pdfs, tmp_path):
    res = convert_file(
        pdfs["report"], "docx", out_dir=tmp_path, overwrite=True,
        overrides={"review_marks": False},
    )
    d = _docx(res.outputs[0].path)
    head = [p.text for p in d.paragraphs[:3]]
    assert not any("自动转换" in t for t in head)
    # 关掉文档内标注不影响 CLI 告警——两者是不同渠道
    assert any("推断" in w for w in res.warnings)


def test_scanned_pdf_rejected_with_actionable_message(pdfs, tmp_path):
    """扫描件转 Word 只会得到一堆图片。提前拦住并给出替代方案。"""
    with pytest.raises(ConversionError) as ei:
        convert_file(pdfs["scanned"], "docx", out_dir=tmp_path, overwrite=True)

    msg = str(ei.value)
    assert "没有文本层" in msg
    assert "OCR" in msg
    assert "-t png" in msg, "应给出可执行的替代方案"


def test_page_selection_in_word(pdfs, tmp_path):
    res = convert_file(
        pdfs["report"], "docx", out_dir=tmp_path, overwrite=True,
        overrides={"pages": "1"},
    )
    d = _docx(res.outputs[0].path)
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Quarterly Report" in text
    # 第 2 页的两栏内容不该出现
    assert "Two column section" not in text


# ── PDF 结构操作：真正无损 ──


def test_page_selection_is_lossless(pdfs, tmp_path):
    """改对象图而不重新渲染，所以是 LOSSLESS 而非 LOSSY。

    和「转成图片再拼回 PDF」有本质区别——后者会把文字变成像素。
    """
    res = convert_file(
        pdfs["plain"], "pdf", out_dir=tmp_path, overwrite=True,
        overrides={"pages": "1,3"},
    )
    assert res.loss is LossTier.LOSSLESS

    import pypdf

    r = pypdf.PdfReader(str(res.outputs[0].path))
    assert len(r.pages) == 2
    # 文字必须还在（没被栅格化）
    assert "Page 1" in r.pages[0].extract_text()
    assert "Page 3" in r.pages[1].extract_text()


def test_split_produces_one_file_per_page(pdfs, tmp_path):
    res = convert_file(
        pdfs["plain"], "pdf", out_dir=tmp_path, overwrite=True,
        overrides={"split": True},
    )
    assert len(res.outputs) == 3
    names = [o.path.name for o in res.outputs]
    assert len(set(names)) == 3, f"文件名重复：{names}"

    import pypdf

    for o in res.outputs:
        assert len(pypdf.PdfReader(str(o.path)).pages) == 1


def test_rotate_applies(pdfs, tmp_path):
    res = convert_file(
        pdfs["plain"], "pdf", out_dir=tmp_path, overwrite=True,
        overrides={"rotate": 90},
    )
    import pypdf

    r = pypdf.PdfReader(str(res.outputs[0].path))
    assert r.pages[0].get("/Rotate") == 90


def test_rotate_rejects_non_multiple_of_90(pdfs, tmp_path):
    with pytest.raises(ConversionError, match="90 的倍数"):
        convert_file(
            pdfs["plain"], "pdf", out_dir=tmp_path, overwrite=True,
            overrides={"rotate": 45},
        )


def test_merge_pdfs(pdfs, tmp_path):
    from convcore.plugins.document_pdf import merge_pdfs

    dst = tmp_path / "merged.pdf"
    total = merge_pdfs([pdfs["plain"], pdfs["report"]], dst)
    assert total == 5  # plain 3 页 + report 2 页

    import pypdf

    assert len(pypdf.PdfReader(str(dst)).pages) == 5


def test_encryption_without_cryptography_gives_actionable_error(pdfs, tmp_path):
    """缺 cryptography 时必须明确报错，绝不静默降级到 RC4。

    RC4 已被破解。降级会让用户以为文件是加密安全的，比直接失败危险得多。
    """
    pytest.importorskip
    try:
        import cryptography  # noqa: F401

        pytest.skip("cryptography 已装，测不到缺依赖的路径")
    except ImportError:
        pass

    with pytest.raises(ConversionError, match="cryptography"):
        convert_file(
            pdfs["plain"], "pdf", out_dir=tmp_path, overwrite=True,
            overrides={"user_password": "x"},
        )


def test_encryption_roundtrip(pdfs, tmp_path):
    pytest.importorskip("cryptography", reason="AES-256 加密需要 cryptography")
    res = convert_file(
        pdfs["plain"], "pdf", out_dir=tmp_path, overwrite=True,
        overrides={"user_password": "s3cret"},
    )
    import pypdf

    r = pypdf.PdfReader(str(res.outputs[0].path))
    assert r.is_encrypted
    assert r.decrypt("s3cret")
    assert len(r.pages) == 3


# ── PDF → 图片：经双桥接 ──


def test_pdf_to_png_one_file_per_page(pdfs, tmp_path):
    """PDF → [document] → [raster] → PNG：跨两次域桥接。"""
    res = convert_file(pdfs["plain"], "png", out_dir=tmp_path, overwrite=True)

    assert len(res.outputs) == 3
    assert res.loss is LossTier.STRUCTURAL
    assert any("栅格化" in w for w in res.warnings)
    assert "[document]" in res.path_desc and "[raster]" in res.path_desc

    from PIL import Image

    with Image.open(res.outputs[0].path) as im:
        # 150dpi 下 A4（595×842pt）约 1240×1754
        assert 1200 < im.width < 1300


def test_dpi_controls_render_size(pdfs, tmp_path):
    lo = convert_file(
        pdfs["report"], "png", out_dir=tmp_path / "lo", overwrite=True,
        overrides={"pages": "1", "dpi": 72},
    )
    hi = convert_file(
        pdfs["report"], "png", out_dir=tmp_path / "hi", overwrite=True,
        overrides={"pages": "1", "dpi": 300},
    )
    from PIL import Image

    with Image.open(lo.outputs[0].path) as a, Image.open(hi.outputs[0].path) as b:
        assert b.width > a.width * 3.5


def test_dpi_out_of_range_rejected(pdfs, tmp_path):
    with pytest.raises(ConversionError, match="dpi"):
        convert_file(
            pdfs["report"], "png", out_dir=tmp_path, overwrite=True,
            overrides={"dpi": 5000},
        )


# ── 图搜索：只写插件的边必须进图 ──


def test_write_only_plugin_edges_are_in_graph():
    """回归测试：注册表曾按「必须有 parse」筛插件，导致只写插件的边不进图。

    症状是「找不到转换路径」而插件明明已启用，非常难查。pdf2docx 和 pypdf
    都是只写插件。
    """
    from convcore.pipeline import Converter

    conv = Converter()
    assert DOCX in conv.registry.writable(), "pdf2docx 的写能力没进注册表"
    assert "application/pdf" in conv.registry.writable(), "pypdf 的写能力没进注册表"

    path = conv.plan("application/pdf", DOCX)
    assert len(path.edges) == 2
    assert path.loss is LossTier.INFERRED
    # 必须选 pdf2docx 而不是绕 markup 域抽文本——后者更快但丢表格
    assert any(getattr(e.plugin, "name", "") == "pdf2docx" for e in path.edges)


def test_fidelity_beats_speed_within_same_loss_tier():
    """同损失等级内，保真度优先于速度。

    回归测试：加 markup 域后出现了 PDF → [document] → [markup] → DOCX 这条
    更廉价的路径（代价 13 vs 43），两者都是 INFERRED，于是时间代价成了决定因素，
    图搜索选了只抽文本的桥接路径——产出的 Word 没有表格。
    """
    from convcore.pipeline import Converter

    conv = Converter()
    path = conv.plan("application/pdf", DOCX)

    plugins = [getattr(e.plugin, "name", "内置") for e in path.edges]
    assert "pdf2docx" in plugins, f"选错了路径：{path.describe()} via {plugins}"
    assert path.fidelity_penalty == 0.0
    # 桥接路径确实更快，所以不是时间在起作用
    assert path.time_cost > 40
