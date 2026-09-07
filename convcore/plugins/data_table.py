"""表格数据：CSV / TSV / JSON Lines / XLSX / Parquet。

两个必须做对的点，都是数据转换的经典事故来源：

**编码嗅探。** 中文环境的 CSV 大量是 GBK（Excel 默认导出就是），按 UTF-8 硬解
会得到乱码。而且 Excel 生成的 UTF-8 CSV 带 BOM，不处理会让第一个列名多个 ﻿，
后续按列名取值全部失败。

**类型推断的保守原则。** 把 "007" 判成 int 会丢前导零，把 18 位身份证号判成
float 会变成科学计数法且精度丢失。这类损坏是不可逆的，且用户往往在很久之后
才发现。所以宁可留字符串也不猜——见 ir/data.py 的 infer_type。
"""

from __future__ import annotations

import csv
import io
from typing import Any

from ..ir.data import Sheet, Table, coerce, infer_type
from ..policy import Policy, Strategy
from ..registry import register
from ..types import Capability, ConversionError, LossTier, MediaType, Resource

CSV_T = "text/csv"
TSV = "text/tab-separated-values"
JSONL = "application/x-ndjson"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PARQUET = "application/vnd.apache.parquet"

_OPTS = frozenset({"sheet", "delimiter", "infer_types", "header", "encoding"})

# 单表行数上限。超过时报错而不是静默截断——静默丢数据比失败危险得多。
MAX_ROWS = 1_000_000


def _decode_csv(data: bytes, forced: str | None = None) -> tuple[str, str]:
    """解码 CSV，返回 (文本, 实际编码)。

    Excel 导出的 UTF-8 CSV 带 BOM，utf-8-sig 会吃掉它。不处理的话第一个列名
    会变成 "﻿姓名"，之后按 "姓名" 取列全部取不到——很难查的坑。
    """
    if forced:
        return data.decode(forced, "replace"), forced
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig"), "utf-8-sig"
    for enc in ("utf-8", "gb18030", "big5", "shift_jis", "latin-1"):
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace"), "utf-8(replace)"


def _sniff_delimiter(sample: str, default: str = ",") -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return default


def _apply_types(sheet: Sheet, enabled: bool, pol: Policy) -> None:
    """推断并应用列类型。

    默认开启，因为 CSV→JSON 时数字该是数字而不是字符串。但推断本身是保守的，
    且把「保持字符串」的列数记进告警，让用户知道有列没被转换。
    """
    if not sheet.rows:
        return
    if not enabled:
        sheet.types = ["str"] * len(sheet.columns)
        return

    kept: list[str] = []
    for ci in range(len(sheet.columns)):
        col = [r[ci] if ci < len(r) else None for r in sheet.rows]
        kind = infer_type(col)
        sheet.types.append(kind)
        if kind != "str":
            for r in sheet.rows:
                if ci < len(r):
                    r[ci] = coerce(r[ci], kind)
        else:
            kept.append(sheet.columns[ci])

    # 只在「看起来像数字但被判成字符串」时才告警，否则噪音太大
    suspicious = [
        c for c, k in zip(sheet.columns, sheet.types, strict=False)
        if k == "str" and _looks_numeric(sheet, c)
    ]
    if suspicious:
        pol.note(
            "类型保护",
            Strategy.APPROXIMATE,
            f"列 {', '.join(suspicious[:5])} 含前导零或超长数字，"
            f"已保持为文本以免丢失精度",
        )


def _looks_numeric(sheet: Sheet, col_name: str) -> bool:
    try:
        ci = sheet.columns.index(col_name)
    except ValueError:
        return False
    sample = [r[ci] for r in sheet.rows[:20] if ci < len(r) and r[ci] not in (None, "")]
    if not sample:
        return False
    return all(str(v).strip().lstrip("-+").replace(".", "").isdigit() for v in sample)


@register
class TableCodec:
    name = "data-table"
    domain = "table"

    reads = (
        Capability(CSV_T, LossTier.LOSSLESS, 1.0, _OPTS),
        Capability(TSV, LossTier.LOSSLESS, 1.0, _OPTS),
        Capability(JSONL, LossTier.LOSSLESS, 1.0, _OPTS),
        # XLSX 是 STRUCTURAL：公式、单元格格式、图表、合并单元格都不进 IR
        Capability(XLSX, LossTier.STRUCTURAL, 3.0, _OPTS, fidelity=0.8),
    )

    writes = (
        Capability(CSV_T, LossTier.LOSSLESS, 1.0, _OPTS),
        Capability(TSV, LossTier.LOSSLESS, 1.0, _OPTS),
        Capability(JSONL, LossTier.LOSSLESS, 1.0, _OPTS),
        Capability(XLSX, LossTier.LOSSLESS, 3.0, _OPTS),
    )

    def available(self) -> tuple[bool, str]:
        try:
            import openpyxl  # noqa: F401

            return True, ""
        except ImportError:
            return False, "缺少 openpyxl（Excel 读写）：uv pip install openpyxl"

    def parse(self, resources: list[Resource], options: dict[str, Any]) -> Table:
        res = resources[0]
        mt = res.media_type
        pol: Policy = options.get("_policy") or Policy()
        infer = options.get("infer_types", True)

        if mt in (CSV_T, TSV):
            table = self._read_csv(res, mt, options, pol)
        elif mt == JSONL:
            table = self._read_jsonl(res, pol)
        elif mt == XLSX:
            table = self._read_xlsx(res, options, pol)
        elif mt == PARQUET:
            table = self._read_parquet(res, pol)
        else:
            raise ConversionError(f"table 插件不支持读 {mt}")

        for s in table.sheets:
            if len(s.rows) > MAX_ROWS:
                raise ConversionError(
                    f"表 {s.name} 有 {len(s.rows)} 行，超过上限 {MAX_ROWS}。"
                    "大表建议用专门的数据工具处理，避免内存耗尽。"
                )
            _apply_types(s, bool(infer), pol)

        return table

    def _read_csv(
        self, res: Resource, mt: MediaType, options: dict[str, Any], pol: Policy
    ) -> Table:
        text, enc = _decode_csv(res.read_bytes(), options.get("encoding"))
        if enc not in ("utf-8", "utf-8-sig"):
            pol.note(
                "文件编码",
                Strategy.APPROXIMATE,
                f"源文件为 {enc} 编码（非 UTF-8），已按该编码解码；输出统一为 UTF-8",
            )

        delim = options.get("delimiter")
        if not delim:
            delim = "\t" if mt == TSV else _sniff_delimiter(text[:8192])

        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delim)
        rows = [r for r in reader if any(c.strip() for c in r)]
        if not rows:
            return Table(sheets=[Sheet()], source_type=mt)

        has_header = options.get("header", True)
        if has_header:
            columns = [c.strip() or f"列{i+1}" for i, c in enumerate(rows[0])]
            body = rows[1:]
        else:
            columns = [f"列{i+1}" for i in range(len(rows[0]))]
            body = rows

        # 补齐短行，避免后续按索引访问越界
        width = len(columns)
        body = [r + [None] * (width - len(r)) if len(r) < width else r[:width] for r in body]

        return Table(
            sheets=[Sheet(name="Sheet1", columns=columns, rows=body)],
            source_type=mt,
            meta={"encoding": enc, "delimiter": delim},
        )

    def _read_jsonl(self, res: Resource, pol: Policy) -> Table:
        import json

        text, _ = _decode_csv(res.read_bytes())
        records: list[dict[str, Any]] = []
        for n, line in enumerate(text.splitlines(), 1):
            line = line.strip()
            if not line:
                continue
            try:
                obj = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ConversionError(f"JSON Lines 第 {n} 行解析失败：{exc.msg}") from exc
            if not isinstance(obj, dict):
                raise ConversionError(
                    f"JSON Lines 第 {n} 行不是对象（表格要求每行是记录）"
                )
            records.append(obj)

        # 列的并集，保持首次出现顺序
        columns: list[str] = []
        for r in records:
            for k in r:
                if k not in columns:
                    columns.append(k)

        nested = sum(
            1 for r in records for v in r.values() if isinstance(v, (dict, list))
        )
        if nested:
            pol.note(
                "嵌套字段",
                Strategy.EMULATE,
                f"{nested} 个嵌套值已序列化为 JSON 字符串（表格无法表达嵌套）",
            )

        import json as _j

        rows = [
            [
                _j.dumps(r.get(c), ensure_ascii=False)
                if isinstance(r.get(c), (dict, list))
                else r.get(c)
                for c in columns
            ]
            for r in records
        ]
        return Table(sheets=[Sheet(columns=columns, rows=rows)], source_type=JSONL)

    def _read_xlsx(self, res: Resource, options: dict[str, Any], pol: Policy) -> Table:
        import openpyxl

        wb = openpyxl.load_workbook(
            io.BytesIO(res.read_bytes()),
            read_only=True,
            data_only=True,  # 取公式的计算结果而非公式本身
        )
        try:
            want = options.get("sheet")
            names = wb.sheetnames
            if want:
                if want not in names:
                    raise ConversionError(
                        f"工作簿里没有名为 {want!r} 的表。可选：{', '.join(names)}"
                    )
                names = [want]

            sheets: list[Sheet] = []
            for name in names:
                ws = wb[name]
                rows = [
                    [c for c in row]
                    for row in ws.iter_rows(values_only=True)
                    if any(c is not None and str(c).strip() for c in row)
                ]
                if not rows:
                    continue
                columns = [
                    str(c).strip() if c is not None else f"列{i+1}"
                    for i, c in enumerate(rows[0])
                ]
                width = len(columns)
                body = [
                    list(r) + [None] * (width - len(r)) if len(r) < width else list(r[:width])
                    for r in rows[1:]
                ]
                sheets.append(Sheet(name=name, columns=columns, rows=body))

            pol.note(
                "Excel 特性",
                Strategy.DROP,
                "公式已取计算结果；单元格格式、图表、批注和合并单元格未保留",
            )
            if len(wb.sheetnames) > 1 and not want:
                pol.note(
                    "多工作表",
                    Strategy.EMULATE,
                    f"共 {len(wb.sheetnames)} 张表，用 --sheet 指定单张",
                )
            return Table(sheets=sheets or [Sheet()], source_type=XLSX)
        finally:
            wb.close()

    def _read_parquet(self, res: Resource, pol: Policy) -> Table:
        try:
            import pyarrow.parquet as pq
        except ImportError as exc:
            raise ConversionError(
                "读 Parquet 需要 pyarrow（约 29MB）：uv pip install pyarrow"
            ) from exc

        tbl = pq.read_table(io.BytesIO(res.read_bytes()))
        columns = list(tbl.column_names)
        py = tbl.to_pylist()
        rows = [[r.get(c) for c in columns] for r in py]
        return Table(sheets=[Sheet(columns=columns, rows=rows)], source_type=PARQUET)

    # ── 写 ──

    def write(
        self, ir: Table, target: MediaType, options: dict[str, Any]
    ) -> list[Resource]:
        pol: Policy = options.get("_policy") or Policy()

        if target == XLSX:
            return [self._write_xlsx(ir, options, pol)]

        # 单表格式遇到多表：一进多出，每表一个文件。
        # 静默只写第一张是数据丢失。
        if len(ir.sheets) > 1:
            pol.note(
                "多工作表",
                Strategy.EMULATE,
                f"{len(ir.sheets)} 张表已分别输出为独立文件",
            )
            return [
                self._write_flat(s, target, options, pol, name=s.name)
                for s in ir.sheets
            ]

        return [self._write_flat(ir.primary, target, options, pol)]

    def _write_flat(
        self,
        sheet: Sheet,
        target: MediaType,
        options: dict[str, Any],
        pol: Policy,
        name: str | None = None,
    ) -> Resource:
        if target == JSONL:
            import json

            lines = [
                json.dumps(rec, ensure_ascii=False, default=str)
                for rec in sheet.records()
            ]
            data = ("\n".join(lines) + "\n").encode("utf-8")
            return Resource(media_type=target, data=data, name=name)

        delim = options.get("delimiter") or ("\t" if target == TSV else ",")
        buf = io.StringIO(newline="")
        # QUOTE_MINIMAL + 显式 lineterminator：默认的 \r\n 在 Excel 里没问题，
        # 但在 Linux 工具链里会多出 ^M
        w = csv.writer(buf, delimiter=delim, lineterminator="\n")
        w.writerow(sheet.columns)
        for row in sheet.rows:
            w.writerow(["" if v is None else v for v in row])

        # 带 BOM 写出：Excel 打开无 BOM 的 UTF-8 CSV 会把中文显示成乱码。
        # 这是 Windows 用户最常遇到的问题，默认加 BOM 更实用。
        text = buf.getvalue()
        if options.get("bom", True):
            data = b"\xef\xbb\xbf" + text.encode("utf-8")
        else:
            data = text.encode("utf-8")
        return Resource(media_type=target, data=data, name=name)

    def _write_xlsx(
        self, ir: Table, options: dict[str, Any], pol: Policy
    ) -> Resource:
        import openpyxl
        from openpyxl.styles import Font

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        for s in ir.sheets:
            # Excel 表名限制：31 字符，且不能含 : \ / ? * [ ]
            safe = s.name[:31]
            for ch in ':\\/?*[]':
                safe = safe.replace(ch, "_")
            ws = wb.create_sheet(safe or "Sheet1")

            ws.append(s.columns)
            for cell in ws[1]:
                cell.font = Font(bold=True)

            for row in s.rows:
                ws.append(["" if v is None else v for v in row])

            # 冻结表头行，大表浏览时表头始终可见
            ws.freeze_panes = "A2"

        buf = io.BytesIO()
        wb.save(buf)
        return Resource(media_type=XLSX, data=buf.getvalue())


def table_to_doctree(table: Table, opts: dict[str, Any], pol: Policy):
    """表格 → 语义文档树桥接。渲染成 Markdown/HTML 表格。

    用途是把数据变成人看的东西：CSV → Markdown 表格贴进文档，
    或 XLSX → HTML 直接在浏览器里看。
    """
    from ..ir.markup import HEADING, TABLE, Block, Cell as MCell, DocTree, Span
    from ..ir.markup import PARAGRAPH

    blocks: list[Block] = []
    multi = len(table.sheets) > 1

    for s in table.sheets:
        if multi:
            blocks.append(Block(HEADING, spans=[Span(s.name)], level=2))

        rows: list[list[MCell]] = [
            [
                MCell(blocks=[Block(PARAGRAPH, spans=[Span(str(c))])], header=True)
                for c in s.columns
            ]
        ]
        for row in s.rows:
            rows.append(
                [
                    MCell(
                        blocks=[
                            Block(
                                PARAGRAPH,
                                spans=[Span("" if v is None else str(v))],
                            )
                        ]
                    )
                    for v in row
                ]
            )
        blocks.append(Block(TABLE, rows=rows))

    return DocTree(blocks=blocks, source_type=table.source_type)
