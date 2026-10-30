"""扫描 PDF / 图片 OCR。

这不是版面还原：Tesseract 只给文字，段落、标题、表格都要靠行/空行重新分组。
所以它只补「没有文本层」这一类源，并且固定标 INFERRED。

选用 Tesseract 而不是 PaddleOCR/Docling：它体积小（约 50MB）、语言包独立、
离线可跑，且不引入 PyTorch。中文需要额外安装 `chi_sim` 语言数据。
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from ..ir.document import DocumentRef
from ..policy import Policy, Strategy
from ..registry import register
from ..types import Capability, ConversionError, LossTier, MediaType, Resource
from ._scratch import workdir

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MARKDOWN = "text/markdown"
PLAIN = "text/plain"

_IMAGE_TYPES = (
    "image/png",
    "image/jpeg",
    "image/webp",
    "image/tiff",
    "image/bmp",
    "image/gif",
    "image/heic",
    "image/avif",
)

_OPTS = frozenset({"pages", "password", "ocr_lang", "ocr_dpi"})


@dataclass(slots=True)
class OcrText:
    """OCR 的中间表示：按页保存纯文本，不假装有版面结构。"""

    pages: list[str]
    source_type: MediaType
    language: str
    engine: str = "tesseract"


def _tesseract_path():
    from ..engines.probe import ENGINES, probe_one

    spec = next((s for s in ENGINES if s.key == "tesseract"), None)
    if spec is None:
        return None
    info = probe_one(spec)
    return info.path if info.available else None


def _load_pytesseract():
    """配置并返回 pytesseract；未装包装/二进制时返回 (None, 原因)。"""
    try:
        import pytesseract
    except ImportError:
        return None, "缺少 pytesseract：pip install pytesseract"
    path = _tesseract_path()
    if path is None:
        return None, (
            "缺少 Tesseract OCR 可执行文件。装法见 "
            "https://github.com/UB-Mannheim/tesseract/wiki"
        )
    pytesseract.pytesseract.tesseract_cmd = str(path)
    return pytesseract, ""


def _choose_language(pytesseract, requested: str | None) -> str:
    try:
        available = set(pytesseract.get_languages(config=""))
    except Exception:
        available = {"eng"}
    if requested:
        if requested not in available:
            raise ConversionError(
                f"OCR 语言 {requested} 未安装。已安装：{', '.join(sorted(available))}"
            )
        return requested
    return "chi_sim+eng" if {"chi_sim", "eng"} <= available else "eng"


@register
class OcrCodec:
    """扫描 PDF / 图片 → 可编辑文本、Markdown、DOCX。"""

    name = "tesseract"
    domain = "ocr"
    reads = (
        Capability(PDF, LossTier.INFERRED, 20.0, _OPTS, fidelity=0.4),
        *(
            Capability(mt, LossTier.INFERRED, 20.0, _OPTS, fidelity=0.4)
            for mt in _IMAGE_TYPES
        ),
    )
    writes = (
        Capability(DOCX, LossTier.INFERRED, 8.0, _OPTS, fidelity=0.9, opaque=True),
        Capability(MARKDOWN, LossTier.INFERRED, 5.0, _OPTS, fidelity=0.9, opaque=True),
        Capability(PLAIN, LossTier.INFERRED, 5.0, _OPTS, fidelity=0.9, opaque=True),
    )

    def available(self) -> tuple[bool, str]:
        _, reason = _load_pytesseract()
        return (False, reason) if reason else (True, "")

    def parse(self, resources: list[Resource], options: dict[str, Any]) -> OcrText:
        pytesseract, reason = _load_pytesseract()
        if pytesseract is None:
            raise ConversionError(f"OCR 不可用：{reason}")
        if len(resources) != 1:
            raise ConversionError("OCR 一次只处理一个源文件")

        lang = _choose_language(pytesseract, options.get("ocr_lang"))
        dpi = int(options.get("ocr_dpi") or 300)
        if dpi < 72 or dpi > 600:
            raise ConversionError(f"OCR dpi 超出合理范围（72-600）：{dpi}")

        source = resources[0]
        tmp = workdir(options)
        path = source.materialize(tmp)

        if source.media_type == PDF:
            pages = self._ocr_pdf(pytesseract, path, lang, dpi, options, tmp)
        else:
            pages = [self._ocr_image(pytesseract, path, lang)]

        if not any(p.strip() for p in pages):
            raise ConversionError(
                "OCR 没有识别出文字。检查图片清晰度/语言包，"
                "或换更高的 --ocr-dpi 重试"
            )
        return OcrText(
            pages=pages,
            source_type=source.media_type,
            language=lang,
        )

    def _ocr_image(self, pytesseract, path: Path, lang: str) -> str:
        from PIL import Image

        with Image.open(path) as image:
            if getattr(image, "n_frames", 1) > 1:
                image.seek(0)
            return pytesseract.image_to_string(image, lang=lang, config="--psm 6")

    def _ocr_pdf(
        self,
        pytesseract,
        path: Path,
        lang: str,
        dpi: int,
        options: dict[str, Any],
        tmp: Path,
    ) -> list[str]:
        import pymupdf
        from PIL import Image

        from .document import _open_doc, parse_page_spec

        doc = _open_doc(
            Resource(media_type=PDF, path=path),
            tmp,
            options.get("password"),
        )
        try:
            indices = parse_page_spec(options.get("pages"), doc.page_count)
            pages: list[str] = []
            for index in indices:
                pix = doc[index].get_pixmap(dpi=dpi, alpha=False)
                image = Image.frombytes(
                    "RGB", (pix.width, pix.height), pix.samples
                )
                pages.append(
                    pytesseract.image_to_string(
                        image, lang=lang, config="--psm 6"
                    )
                )
            return pages
        finally:
            doc.close()

    def write(
        self, ir: OcrText, target: MediaType, options: dict[str, Any]
    ) -> list[Resource]:
        pol: Policy = options.get("_policy") or Policy()
        if not isinstance(ir, OcrText):
            # 图搜索可能选择「pymupdf READ + OCR WRITE」这条更便宜的路：
            # READ 只给 DocumentRef，OCR 在 write 阶段补做。
            if isinstance(ir, DocumentRef):
                ir = self.parse([ir.source], options)
            else:
                raise ConversionError(f"OCR 收到不支持的中间表示：{type(ir).__name__}")
        pol.note(
            "OCR",
            Strategy.APPROXIMATE,
            f"文字由 Tesseract OCR 识别（语言 {ir.language}），"
            "段落按空行切分；错字、表格和标题层级都需要人工复核",
        )

        if target == DOCX:
            return [Resource(media_type=target, data=self._docx(ir))]
        if target == MARKDOWN:
            return [Resource(media_type=target, data=self._markdown(ir).encode("utf-8"))]
        if target == PLAIN:
            return [Resource(media_type=target, data=self._plain(ir).encode("utf-8"))]
        raise ConversionError(f"OCR 不支持写出 {target}")

    def _docx(self, ir: OcrText) -> bytes:
        import docx
        from docx.shared import Pt, RGBColor

        document = docx.Document()
        note = document.add_paragraph(
            "本文档由 OCR 自动识别生成，文字和排版均可能有误，请对照原件复核。"
        )
        note.runs[0].font.size = Pt(9)
        note.runs[0].font.color.rgb = RGBColor(0x99, 0x6B, 0x00)

        for n, page in enumerate(ir.pages, 1):
            if n > 1:
                document.add_page_break()
            paragraphs = [p.strip() for p in page.split("\n\n")]
            for text in paragraphs:
                if text:
                    document.add_paragraph(text.replace("\n", " "))

        buf = io.BytesIO()
        document.save(buf)
        return buf.getvalue()

    def _markdown(self, ir: OcrText) -> str:
        parts = [
            "> 本文档由 OCR 自动识别生成，文字和排版均可能有误，请对照原件复核。",
            "",
        ]
        for n, page in enumerate(ir.pages, 1):
            if n > 1:
                parts.append("---")
                parts.append("")
            parts.append(f"<!-- page {n} -->")
            parts.append("")
            parts.extend(p.strip() for p in page.split("\n\n") if p.strip())
            parts.append("")
        return "\n".join(parts)

    def _plain(self, ir: OcrText) -> str:
        return "\n\n".join(
            page.strip() for page in ir.pages if page.strip()
        ) + "\n"
