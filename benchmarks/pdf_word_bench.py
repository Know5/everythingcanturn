"""PDF→Word capability benchmark.

Run from the repository root:

    python benchmarks/fetch_samples.py
    python benchmarks/pdf_word_bench.py

The script only writes under `benchmarks/work/` (inside the project). It
measures text recall, headings, tables, images, text boxes, runtime and
structural size for each sample, and writes `report.json` + `report.md`.
"""

from __future__ import annotations

import json
import re
import sys
import time
import zipfile
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent))

from convcore import ConversionError, convert_file  # noqa: E402

SAMPLES = ROOT / "samples"
WORK = ROOT / "work"

SOURCES = {
    "attention": SAMPLES / "attention.pdf",
    "fw9": SAMPLES / "fw9.pdf",
    "scanned": SAMPLES / "scanned.pdf",
}


def tokens(text: str) -> Counter[str]:
    return Counter(re.findall(r"[A-Za-z0-9]+|[\u4e00-\u9fff]", text.lower()))


def docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def docx_stats(path: Path) -> dict:
    import docx

    document = docx.Document(str(path))
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8", "replace")
        images = [
            n for n in archive.namelist() if n.startswith("word/media/")
        ]
    return {
        "paragraphs": len(document.paragraphs),
        "headings": sum(
            1 for p in document.paragraphs if p.style.name.startswith("Heading")
        ),
        "tables": len(document.tables),
        "images": len(images),
        "textboxes": xml.count("txbxContent"),
        "bytes": path.stat().st_size,
    }


def pdf_stats(path: Path) -> dict:
    import pymupdf

    with pymupdf.open(path) as doc:
        text = "\n".join(page.get_text("text") for page in doc)
        images = sum(len(page.get_images(full=False)) for page in doc)
        tables = 0
        for page in doc:
            try:
                tables += len(page.find_tables().tables)
            except Exception:
                pass
        return {
            "pages": doc.page_count,
            "text_chars": len(text.strip()),
            "images": images,
            "tables": tables,
            "text": text,
        }


def run_one(name: str, pdf: Path, out_root: Path) -> dict:
    out_dir = out_root / name
    out_dir.mkdir(parents=True, exist_ok=True)
    info = pdf_stats(pdf)
    result = {
        "sample": name,
        "file": pdf.name,
        "pdf": {k: v for k, v in info.items() if k != "text"},
        "status": None,
    }
    started = time.perf_counter()
    try:
        res = convert_file(
            pdf,
            "docx",
            out_dir=out_dir,
            overwrite=True,
            overrides={"review_marks": False},
        )
    except ConversionError as exc:
        result.update(
            status="rejected",
            elapsed=round(time.perf_counter() - started, 2),
            error=str(exc),
        )
        return result
    except Exception as exc:
        result.update(
            status="error",
            elapsed=round(time.perf_counter() - started, 2),
            error=f"{type(exc).__name__}: {exc}",
        )
        return result

    elapsed = round(time.perf_counter() - started, 2)
    docx_path = res.outputs[0].path
    docx_info = docx_stats(docx_path)
    docx_tokens = tokens(docx_text(docx_path))
    pdf_tokens = tokens(info["text"])
    matched = sum(min(n, docx_tokens.get(t, 0)) for t, n in pdf_tokens.items())
    recall = matched / sum(pdf_tokens.values()) if pdf_tokens else None
    result.update(
        status="ok",
        elapsed=elapsed,
        loss=res.loss.name,
        warnings=res.warnings,
        docx=docx_info,
        text_recall=round(recall, 4) if recall is not None else None,
        docx_tokens=sum(docx_tokens.values()),
        pdf_tokens=sum(pdf_tokens.values()),
    )
    return result


def main() -> int:
    missing = [str(p) for p in SOURCES.values() if not p.is_file()]
    if missing:
        print("missing samples; run benchmarks/fetch_samples.py:", missing, file=sys.stderr)
        return 2
    WORK.mkdir(parents=True, exist_ok=True)
    results = []
    for name, pdf in SOURCES.items():
        print(f"== {name} ==", flush=True)
        result = run_one(name, pdf, WORK)
        results.append(result)
        print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)

    (WORK / "report.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    lines = [
        "# PDF→Word benchmark",
        "",
        "| sample | status | pages | recall | headings | tables | images | textboxes | secs |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for r in results:
        d = r.get("docx", {})
        lines.append(
            "| {sample} | {status} | {pages} | {recall} | {headings} | {tables} | "
            "{images} | {textboxes} | {elapsed} |".format(
                sample=r["sample"],
                status=r["status"],
                pages=r["pdf"]["pages"],
                recall="—" if r.get("text_recall") is None else r["text_recall"],
                headings=d.get("headings", "—"),
                tables=d.get("tables", "—"),
                images=d.get("images", "—"),
                textboxes=d.get("textboxes", "—"),
                elapsed=r.get("elapsed", "—"),
            )
        )
    report = "\n".join(lines) + "\n"
    (WORK / "report.md").write_text(report, encoding="utf-8")
    print("\n" + report)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
